"""Threat report generation in PDF, DOCX, XLSX, CSV and JSON."""
import csv
import io
import json
from datetime import datetime, time, timedelta
from xml.sax.saxutils import escape

from django.core.files.base import ContentFile
from django.core.serializers.json import DjangoJSONEncoder
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.utils.text import slugify

from ..models import Alert, Asset, ExternalExposure, Incident, Indicator, Report, Severity
from .analytics import summary
from .compliance import framework_summary
from .exports import safe_cell

MAX_REPORT_INDICATORS = 500
INDICATOR_COLUMNS = ["Value", "Type", "Threat", "Severity", "Risk", "Status", "Source", "Last seen"]


def _parse_day(value, end=False):
    if not value:
        return None
    day = value if hasattr(value, "year") else parse_date(str(value))
    if day is None:
        raise ValueError(f"Invalid date {value!r}; use YYYY-MM-DD.")
    return timezone.make_aware(datetime.combine(day, time.max if end else time.min))


def build_report_data(parameters):
    start = _parse_day(parameters.get("date_from"))
    end = _parse_day(parameters.get("date_to"), end=True)
    if start and end and start > end:
        raise ValueError("date_from must be on or before date_to.")
    min_severity = int(parameters.get("min_severity") or Severity.LOW)

    indicators = Indicator.objects.select_related("source").filter(severity__gte=min_severity)
    alerts = Alert.objects.select_related("indicator", "rule", "asset").filter(severity__gte=min_severity)
    incidents = Incident.objects.filter(severity__gte=min_severity)
    exposures = ExternalExposure.objects.filter(severity__gte=min_severity)
    if start:
        indicators = indicators.filter(last_seen__gte=start)
        alerts = alerts.filter(created_at__gte=start)
        incidents = incidents.filter(last_activity__gte=start)
        exposures = exposures.filter(observed_at__gte=start)
    if end:
        indicators = indicators.filter(first_seen__lte=end)
        alerts = alerts.filter(created_at__lte=end)
        incidents = incidents.filter(created_at__lte=end)
        exposures = exposures.filter(observed_at__lte=end)
    if parameters.get("threat_types"):
        indicators = indicators.filter(threat_type__in=parameters["threat_types"])

    days = 30
    if start:
        days = max(min(((end or timezone.now()) - start).days + 1, 365), 1)
    stats = summary(days=days, now=end if end and end < timezone.now() else None)
    return {
        "title": parameters.get("title") or "Threat Intelligence Report",
        "generated_at": timezone.now(),
        "period": {"from": start.date().isoformat() if start else None, "to": end.date().isoformat() if end else None},
        "min_severity": Severity(min_severity).label,
        "summary": stats,
        "indicators": list(indicators.order_by("-risk_score", "-last_seen")[:MAX_REPORT_INDICATORS]),
        "indicator_count": indicators.count(),
        "alerts": list(alerts.order_by("-created_at")[:200]),
        "alert_count": alerts.count(),
        "incidents": list(incidents.order_by("-priority_score")[:50]),
        "exposures": list(exposures.order_by("-severity", "-observed_at")[:50]),
        "shadow_it": Asset.objects.filter(managed=False).count(),
        "isolated": Asset.objects.filter(is_isolated=True).count(),
        "compliance": framework_summary(),
    }


def _indicator_cells(ind):
    return [
        ind.value, ind.get_ioc_type_display(), ind.get_threat_type_display(), ind.get_severity_display(),
        ind.risk_score, ind.get_status_display(), ind.source.name if ind.source_id else "",
        ind.last_seen.strftime("%Y-%m-%d %H:%M"),
    ]


def _key_metrics(data):
    totals = data["summary"]["totals"]
    return [
        ("Threat posture", f"{data['summary']['posture']['score']} ({data['summary']['posture']['level']})"),
        ("Indicators in scope", data["indicator_count"]),
        ("Actionable indicators", totals["actionable"]),
        ("Critical active indicators", totals["critical_active"]),
        ("High-risk indicators (score >= 80)", totals["high_risk"]),
        ("Average risk score", totals["average_risk"]),
        ("Alerts in scope", data["alert_count"]),
        ("Open alerts", totals["open_alerts"]),
        ("Open incidents in scope", sum(1 for i in data["incidents"] if i.status != Incident.Status.CLOSED)),
        ("Shadow-IT devices", data["shadow_it"]),
        ("Isolated devices", data["isolated"]),
        ("Mean time to resolve (hours)", totals["mttr_hours"] if totals["mttr_hours"] is not None else "n/a"),
    ]


def render_pdf(data):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm, title=data["title"])
    styles = getSampleStyleSheet()
    cell = styles["BodyText"].clone("cell", fontSize=8, leading=10)
    header_bg = colors.HexColor("#1f2937")

    def table(rows, widths=None):
        wrapped = [[Paragraph(escape(str(v)), cell) for v in row] for row in rows]
        t = Table(wrapped, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), header_bg),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#9ca3af")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f4f6")]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        return t

    period = f"{data['period']['from'] or 'beginning'} to {data['period']['to'] or 'now'}"
    story = [
        Paragraph(escape(data["title"]), styles["Title"]),
        Paragraph(escape(f"Generated {data['generated_at']:%Y-%m-%d %H:%M %Z} | Period: {period} | "
                         f"Minimum severity: {data['min_severity']}"), styles["Normal"]),
        Spacer(1, 6 * mm),
        Paragraph("Executive summary", styles["Heading2"]),
        table([["Metric", "Value"], *_key_metrics(data)], widths=[90 * mm, 50 * mm]),
        Spacer(1, 5 * mm),
    ]
    if data["compliance"]:
        story += [Paragraph("Compliance readiness", styles["Heading2"]),
                  table([["Framework", "Readiness", "Controls"]] +
                        [[f["label"], f"{f['readiness']}%", f["controls"]] for f in data["compliance"]],
                        widths=[80 * mm, 30 * mm, 30 * mm]), Spacer(1, 5 * mm)]
    stats = data["summary"]
    story += [
        Paragraph("Severity distribution (actionable)", styles["Heading2"]),
        table([["Severity", "Count"]] + [[r["label"], r["count"]] for r in stats["by_severity"]],
              widths=[60 * mm, 30 * mm]),
        Spacer(1, 5 * mm),
    ]
    if stats["anomalies"]:
        story += [Paragraph("Detected volume anomalies", styles["Heading2"]),
                  table([["Date", "Indicators", "Baseline", "Z-score"]] +
                        [[a["date"], a["count"], a["baseline"], a["zscore"] or "n/a"] for a in stats["anomalies"]]),
                  Spacer(1, 5 * mm)]
    if data["incidents"]:
        story += [Paragraph("Incidents", styles["Heading2"]),
                  table([["Incident", "Severity", "Status", "Priority", "Alerts", "Last activity"]] +
                        [[i.title, i.get_severity_display(), i.get_status_display(), i.priority_score, i.alert_count,
                          i.last_activity.strftime("%Y-%m-%d %H:%M") if i.last_activity else ""]
                         for i in data["incidents"]],
                        widths=[110 * mm, 22 * mm, 25 * mm, 20 * mm, 18 * mm, 32 * mm]),
                  Spacer(1, 5 * mm)]
    if data["exposures"]:
        story += [Paragraph("External exposure", styles["Heading2"]),
                  table([["Finding", "Kind", "Severity", "Status", "Observed"]] +
                        [[e.title, e.get_kind_display(), e.get_severity_display(), e.get_status_display(),
                          e.observed_at.strftime("%Y-%m-%d")] for e in data["exposures"]],
                        widths=[120 * mm, 35 * mm, 22 * mm, 25 * mm, 25 * mm])]
    story += [
        PageBreak(),
        Paragraph(f"Top indicators ({len(data['indicators'])} of {data['indicator_count']})", styles["Heading2"]),
        table([INDICATOR_COLUMNS] + [_indicator_cells(i) for i in data["indicators"]],
              widths=[90 * mm, 22 * mm, 30 * mm, 18 * mm, 12 * mm, 24 * mm, 35 * mm, 30 * mm]),
    ]
    doc.build(story)
    return buffer.getvalue()


def render_docx(data):
    from docx import Document
    from docx.shared import Pt

    document = Document()
    document.add_heading(data["title"], level=0)
    period = f"{data['period']['from'] or 'beginning'} to {data['period']['to'] or 'now'}"
    document.add_paragraph(f"Generated {data['generated_at']:%Y-%m-%d %H:%M %Z}. Period: {period}. "
                           f"Minimum severity: {data['min_severity']}.")

    def add_table(header, rows):
        t = document.add_table(rows=1, cols=len(header))
        t.style = "Light Grid Accent 1"
        for i, text in enumerate(header):
            t.rows[0].cells[i].text = str(text)
        for row in rows:
            cells = t.add_row().cells
            for i, value in enumerate(row):
                cells[i].text = str(value)
                for paragraph in cells[i].paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(8)

    document.add_heading("Executive summary", level=1)
    add_table(["Metric", "Value"], _key_metrics(data))
    if data["compliance"]:
        document.add_heading("Compliance readiness", level=1)
        add_table(["Framework", "Readiness", "Controls"],
                  [[f["label"], f"{f['readiness']}%", f["controls"]] for f in data["compliance"]])
    document.add_heading("Severity distribution", level=1)
    add_table(["Severity", "Count"], [[r["label"], r["count"]] for r in data["summary"]["by_severity"]])
    if data["incidents"]:
        document.add_heading("Incidents", level=1)
        add_table(["Incident", "Severity", "Status", "Priority", "Alerts"],
                  [[i.title, i.get_severity_display(), i.get_status_display(), i.priority_score, i.alert_count]
                   for i in data["incidents"]])
    if data["exposures"]:
        document.add_heading("External exposure", level=1)
        add_table(["Finding", "Kind", "Severity", "Status"],
                  [[e.title, e.get_kind_display(), e.get_severity_display(), e.get_status_display()]
                   for e in data["exposures"]])
    document.add_heading(f"Top indicators ({len(data['indicators'])} of {data['indicator_count']})", level=1)
    add_table(INDICATOR_COLUMNS, [_indicator_cells(i) for i in data["indicators"]])
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_xlsx(data):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    bold = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F2937")

    def sheet(ws, header, rows):
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = bold, fill
        for row in rows:
            ws.append([safe_cell(v) for v in row])
        for column in ws.columns:
            width = max(len(str(c.value or "")) for c in column)
            ws.column_dimensions[column[0].column_letter].width = min(max(width + 2, 10), 60)
        ws.freeze_panes = "A2"

    summary_ws = workbook.active
    summary_ws.title = "Summary"
    sheet(summary_ws, ["Metric", "Value"], _key_metrics(data))
    sheet(workbook.create_sheet("Indicators"), INDICATOR_COLUMNS + ["Tags", "MITRE ATT&CK"],
          [_indicator_cells(i) + [", ".join(i.tags), ", ".join(i.mitre_techniques)] for i in data["indicators"]])
    sheet(workbook.create_sheet("Incidents"), ["Incident", "Severity", "Status", "Priority", "Alerts", "Techniques"],
          [[i.title, i.get_severity_display(), i.get_status_display(), i.priority_score, i.alert_count,
            ", ".join(i.mitre_techniques)] for i in data["incidents"]])
    sheet(workbook.create_sheet("Alerts"), ["Created", "Title", "Severity", "Status", "Detection"],
          [[a.created_at.strftime("%Y-%m-%d %H:%M"), a.title, a.get_severity_display(), a.get_status_display(),
            a.get_detection_type_display()] for a in data["alerts"]])
    sheet(workbook.create_sheet("Exposure"), ["Finding", "Kind", "Severity", "Status", "Observed"],
          [[e.title, e.get_kind_display(), e.get_severity_display(), e.get_status_display(),
            e.observed_at.strftime("%Y-%m-%d")] for e in data["exposures"]])
    sheet(workbook.create_sheet("Compliance"), ["Framework", "Readiness %", "Controls"],
          [[f["label"], f["readiness"], f["controls"]] for f in data["compliance"]])
    sheet(workbook.create_sheet("Threat types"), ["Threat type", "Count"],
          [[r["label"], r["count"]] for r in data["summary"]["by_threat_type"]])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render_csv(data):
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(INDICATOR_COLUMNS)
    for indicator in data["indicators"]:
        writer.writerow([safe_cell(v) for v in _indicator_cells(indicator)])
    return buffer.getvalue().encode("utf-8")


def render_json(data):
    payload = {
        "title": data["title"],
        "generated_at": data["generated_at"],
        "period": data["period"],
        "min_severity": data["min_severity"],
        "summary": data["summary"],
        "compliance": data["compliance"],
        "incidents": [{"id": i.pk, "title": i.title, "severity": i.severity, "status": i.status,
                       "priority": i.priority_score, "alerts": i.alert_count} for i in data["incidents"]],
        "exposures": [{"id": e.pk, "title": e.title, "kind": e.kind, "severity": e.severity, "status": e.status}
                      for e in data["exposures"]],
        "indicators": [
            {"value": i.value, "ioc_type": i.ioc_type, "threat_type": i.threat_type, "severity": i.severity,
             "risk_score": i.risk_score, "status": i.status, "tags": i.tags, "mitre_techniques": i.mitre_techniques,
             "first_seen": i.first_seen, "last_seen": i.last_seen}
            for i in data["indicators"]
        ],
    }
    return json.dumps(payload, cls=DjangoJSONEncoder, indent=2).encode("utf-8")


RENDERERS = {
    Report.Format.PDF: render_pdf,
    Report.Format.DOCX: render_docx,
    Report.Format.XLSX: render_xlsx,
    Report.Format.CSV: render_csv,
    Report.Format.JSON: render_json,
}
CONTENT_TYPES = {
    Report.Format.PDF: "application/pdf",
    Report.Format.DOCX: "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    Report.Format.XLSX: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    Report.Format.CSV: "text/csv",
    Report.Format.JSON: "application/json",
}


def generate_report(report):
    data = build_report_data({**report.parameters, "title": report.name})
    content = RENDERERS[report.file_format](data)
    stamp = timezone.now().strftime("%Y%m%d-%H%M%S")
    filename = f"{slugify(report.name) or 'report'}-{stamp}.{report.file_format}"
    report.summary = {
        "indicator_count": data["indicator_count"],
        "alert_count": data["alert_count"],
        "incident_count": len(data["incidents"]),
        "posture": data["summary"]["posture"],
        "critical_active": data["summary"]["totals"]["critical_active"],
        "size_bytes": len(content),
    }
    report.report_file.save(filename, ContentFile(content), save=False)
    report.save()
    return report


def default_period(days=30):
    today = timezone.localdate()
    return {"date_from": (today - timedelta(days=days - 1)).isoformat(), "date_to": today.isoformat()}
